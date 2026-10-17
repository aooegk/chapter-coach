#!/usr/bin/env python3
"""抽取 .pptx 的正文、表格与备注页，输出 Markdown 和 meta.json。

用法:
    python extract_pptx.py ch03.pptx --chapter-id ch03 --out .ppt-tutor

输出:
    <out>/extracted/ch03.md
    <out>/extracted/ch03.meta.json

依赖: python-pptx (pip install python-pptx)
无 python-pptx 时请走 SKILL.md 的降级链（直接读 XML / LibreOffice 转 PDF）。
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def shape_texts(shape, out: list[str]) -> bool:
    """递归收集形状文本，返回该形状是否包含图片。"""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    has_picture = False
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child in shape.shapes:
            has_picture = shape_texts(child, out) or has_picture
        return has_picture

    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
        has_picture = True

    if getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text.strip()
        if text:
            out.append(text)

    if getattr(shape, "has_table", False):
        table = shape.table
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip().replace("|", "\\|") for cell in row.cells])
        if rows:
            out.append("TABLE::" + json.dumps(rows, ensure_ascii=False))

    return has_picture


def slide_title(slide) -> str:
    try:
        if slide.shapes.title is not None:
            return slide.shapes.title.text.strip()
    except Exception:
        pass
    return ""


def notes_text(slide) -> str:
    try:
        if slide.has_notes_slide:
            return slide.notes_slide.notes_text_frame.text.strip()
    except Exception:
        pass
    return ""


def table_to_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def main() -> int:
    parser = argparse.ArgumentParser(description="抽取 PPTX 文本与备注")
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--chapter-id", required=True, help="如 ch03")
    parser.add_argument("--out", type=Path, default=Path("."), help="课程根目录或 .ppt-tutor 目录")
    args = parser.parse_args()

    if not args.pptx.exists():
        raise SystemExit(f"[error] 文件不存在: {args.pptx}")

    try:
        from pptx import Presentation
    except ImportError:
        print(
            "[error] 未安装 python-pptx。\n"
            "        A 档: pip install python-pptx\n"
            "        B 档: 解压 pptx 读取 ppt/slides/*.xml 与 ppt/notesSlides/*.xml，"
            "或 LibreOffice 转 PDF 后抽文本。",
            file=sys.stderr,
        )
        return 3

    prs = Presentation(str(args.pptx))
    slide_md: list[str] = []
    image_only: list[int] = []
    has_notes = False
    total_chars = 0

    for index, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        has_picture = False
        title_shape = slide.shapes.title
        title_shape_id = getattr(title_shape, "shape_id", None)
        for shape in slide.shapes:
            # 标题已单独输出，避免正文重复一遍
            if title_shape_id is not None and getattr(shape, "shape_id", None) == title_shape_id:
                continue
            has_picture = shape_texts(shape, parts) or has_picture

        title = slide_title(slide)
        body_texts = [p for p in parts if not p.startswith("TABLE::")]
        text_len = sum(len(p) for p in body_texts)

        if text_len < 15 and has_picture:
            image_only.append(index)

        slide_md.append(f"## Slide {index}")
        if title:
            slide_md.append(f"**{title}**")
            slide_md.append("")
        for part in parts:
            if part.startswith("TABLE::"):
                rows = json.loads(part[len("TABLE::"):])
                slide_md.append(table_to_markdown(rows))
            else:
                slide_md.append(part)
            slide_md.append("")
        if not parts:
            slide_md.append("_（本页无可抽取文本）_")
            slide_md.append("")

        note = notes_text(slide)
        if note:
            has_notes = True
            slide_md.append("### 备注")
            slide_md.append(note)
            slide_md.append("")

        total_chars += text_len + len(note)

    title_text = ""
    try:
        first_slide = next(iter(prs.slides), None)
        if first_slide is not None:
            title_text = slide_title(first_slide)
    except Exception:
        pass
    if not title_text:
        try:
            title_text = prs.core_properties.title or ""
        except Exception:
            pass

    out_root = args.out
    extracted = out_root / "extracted"
    extracted.mkdir(parents=True, exist_ok=True)
    md_path = extracted / f"{args.chapter_id}.md"
    meta_path = extracted / f"{args.chapter_id}.meta.json"

    header = [
        f"# {title_text or args.chapter_id}",
        "",
        f"> 来源: {args.pptx.name} | 页数: {len(prs.slides)} | "
        f"抽取时间: {datetime.now(timezone.utc).isoformat()}",
        "",
    ]
    md_path.write_text("\n".join(header + slide_md), encoding="utf-8")

    meta = {
        "chapter_id": args.chapter_id,
        "file": args.pptx.name,
        "sha256": sha256_file(args.pptx),
        "slides": len(prs.slides),
        "has_notes": has_notes,
        "image_only_slides": image_only,
        "text_chars": total_chars,
        "extracted_at": datetime.now(timezone.utc).isoformat(),
        "extractor": "python-pptx",
    }
    meta_path.write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"[ok] 已写出 {md_path}")
    print(f"[ok] 已写出 {meta_path}")
    if image_only:
        print(f"[warn] 疑似图片页（需视觉模型或人工处理）: {image_only}")
    if not has_notes:
        print("[warn] 未发现备注页内容；若课件以备注为主，请确认抽取方式")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
